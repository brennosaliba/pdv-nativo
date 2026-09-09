import re, io
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

DOC = r"C:\Users\Waz\pdv-nativo\docs\HOMOLOGACAO_58_PASSOS.md"
SAIDA = r"C:\temp\Homologacao_TEF_AmericanDay.xlsx"

doc = io.open(DOC, encoding="utf-8").read()
blocos = re.split(r"\n### Passo ", doc)[1:]

OBR = {n: "SIM" for n in range(1, 59)}
for n in [9, 13, 14, 15, 19, 20, 22, 24, 25, 38, 40]:
    OBR[n] = "OPCIONAL"
for n in [41, 42]:
    OBR[n] = "AUTOATENDIMENTO"
for n in [49, 50, 51, 52, 53]:
    OBR[n] = "CONTROLPAY"
OBR[58] = "C6PAY ANDROID"

# O que ja foi feito hoje, 09/09/2026, no sandbox PdC 115998.
FEITO = {1: "aprovado", 6: "aprovado"}


def campo(b, nome):
    m = re.search(r"\*\*" + nome + r":\*\*\s*(.+?)(?=\n\*\*|\n### |\Z)", b, re.S)
    return re.sub(r"\s+", " ", m.group(1)).strip() if m else ""


passos = []
for b in blocos:
    num = int(b[:2])
    titulo = b.split("\n", 1)[0][4:].strip()
    faz = campo(b, "O que você faz")
    conf = campo(b, "O que conferir")
    situ = campo(b, "Situação")
    mv = re.search(r"R\$ ?([\d.]+,\d\d)", titulo) or re.search(r"R\$ ?([\d.]+,\d\d)", faz)
    passos.append({
        "n": num, "titulo": titulo, "obr": OBR[num],
        "valor": ("R$ " + mv.group(1)) if mv else "",
        "situacao": situ, "faz": faz, "conf": conf,
    })

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Roteiro"

TITULO = Font(bold=True, size=11, color="FFFFFF")
FUNDO = PatternFill("solid", fgColor="4A148C")
BORDA = Border(*[Side(style="thin", color="D0D0D0")] * 4)

cols = ["Passo", "Obrigatoriedade", "Valor exato", "O que fazer no caixa",
        "O que conferir", "REQNUM", "Resultado", "Situacao levantada"]
larguras = [8, 16, 14, 62, 62, 16, 14, 20]
ws.append(cols)
for i, (c, w) in enumerate(zip(cols, larguras), start=1):
    ws.cell(row=1, column=i).font = TITULO
    ws.cell(row=1, column=i).fill = FUNDO
    ws.cell(row=1, column=i).alignment = Alignment(vertical="center", wrap_text=True)
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
ws.freeze_panes = "A2"
ws.row_dimensions[1].height = 30

VERDE = PatternFill("solid", fgColor="E8F5E9")
CINZA = PatternFill("solid", fgColor="F5F5F5")
AMARELO = PatternFill("solid", fgColor="FFF8E1")

linha = 2
for p in passos:
    nosso = p["obr"] in ("SIM", "OPCIONAL")
    ws.append([
        f"Passo {p['n']}", p["obr"], p["valor"], p["faz"], p["conf"],
        "", FEITO.get(p["n"], ""), p["situacao"],
    ])
    for col in range(1, len(cols) + 1):
        cel = ws.cell(row=linha, column=col)
        cel.alignment = Alignment(vertical="top", wrap_text=True)
        cel.border = BORDA
        if not nosso:
            cel.fill = CINZA
        elif p["n"] in FEITO:
            cel.fill = VERDE
        elif p["obr"] == "SIM":
            cel.fill = AMARELO
    ws.row_dimensions[linha].height = 46
    linha += 1

# Aba de instrucoes: o que a planilha oficial exige e o que ficou de fora.
wi = wb.create_sheet("Leia primeiro")
wi.column_dimensions["A"].width = 110
for txt, negrito in [
    ("Homologacao TEF PayGo - American Day", True),
    ("", False),
    ("Ponto de captura 115998, CNPJ 62.177.839/0001-57, ambiente de TESTE.", False),
    ("Biblioteca PGWebLib 4.1.50.24, via PayGo Windows 5.1.50.24.", False),
    ("", False),
    ("O QUE A PLANILHA OFICIAL EXIGE", True),
    ("Na coluna 'Retorno do teste', para integracao por biblioteca Windows (DLL),", False),
    ("vai o PWINFO_REQNUM da transacao. E a coluna REQNUM desta planilha.", False),
    ("", False),
    ("CORES", True),
    ("Verde: ja feito e aprovado no sandbox.", False),
    ("Amarelo: obrigatorio, ainda por fazer.", False),
    ("Cinza: nao se aplica a esta integracao (ControlPay, autoatendimento, Android).", False),
    ("", False),
    ("O QUE FICOU DE FORA, E POR QUE", True),
    ("Passos 49 a 53: sao de ControlPay, que e outra integracao.", False),
    ("Passos 41 e 42: autoatendimento, nao e o caso da loja.", False),
    ("Passo 58: C6PAY Android.", False),
    ("", False),
    ("ATENCAO NO AMBIENTE DE TESTE", True),
    ("O campo de redes do PDV precisa ficar VAZIO no sandbox. Com C6PAY forcada,", False),
    ("o terminal responde [NA A116] SERVICO NAO HABILITADO, porque C6PAY e rede", False),
    ("de producao. Ao voltar para producao, recolocar C6PAY, REDE, PIX C6 BANK.", False),
    ("", False),
    ("A porta do pinpad nesta maquina e a COM3 (Gertec PIN Pad PPC), nao a COM5.", False),
]:
    c = wi.cell(row=wi.max_row + 1 if wi.max_row > 1 or wi["A1"].value else 1, column=1, value=txt)
    if negrito:
        c.font = Font(bold=True, size=12)

wb.save(SAIDA)

nossos = [p for p in passos if p["obr"] in ("SIM", "OPCIONAL")]
obrig = [p for p in passos if p["obr"] == "SIM"]
print("gerado:", SAIDA)
print(f"  {len(passos)} passos no total")
print(f"  {len(nossos)} se aplicam a esta integracao")
print(f"  {len(obrig)} obrigatorios, {len(FEITO)} ja aprovados")
print(f"  {len([p for p in nossos if p['valor']])} com valor exato definido pelo roteiro")
